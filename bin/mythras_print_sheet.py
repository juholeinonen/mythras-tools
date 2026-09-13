#!/usr/bin/env python3
"""
Printable Mythras Character Creation Sheets

The two generators beside this script build sheets that calculate as you type. These are
the paper version: every value is an empty box to fill in by hand, and the page is set up to
print on one sheet of A4. Only static information is printed - skill names and the
characteristics each base is worked out from.

Usage:
    python bin/mythras_print_sheet.py              # both variants
    python bin/mythras_print_sheet.py d20 -o out   # one variant, into ./out

Requirements: pip install openpyxl
"""

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

# The same skills as the interactive generators, with the base written out for a person
# to work out rather than as a formula.
STANDARD_SKILLS = [
    ("Athletics", "STR+DEX"),
    ("Boating", "STR+CON"),
    ("Brawn", "STR+SIZ"),
    ("Conceal", "DEX+POW"),
    ("Customs", "INT×2+40"),
    ("Dance", "DEX+CHA"),
    ("Deceit", "INT+CHA"),
    ("Drive", "DEX+POW"),
    ("Endurance", "CON×2"),
    ("Evade", "DEX×2"),
    ("First Aid", "INT+DEX"),
    ("Influence", "CHA×2"),
    ("Insight", "INT+POW"),
    ("Locale", "INT×2"),
    ("Native Tongue", "INT+CHA+40"),
    ("Perception", "INT+POW"),
    ("Ride", "DEX+POW"),
    ("Sing", "CHA+POW"),
    ("Stealth", "DEX+INT"),
    ("Swim", "STR+CON"),
    ("Unarmed", "STR+DEX"),
    ("Willpower", "POW×2"),
    ("Combat Style", "STR+DEX"),
]

CHARACTERISTICS = ["STR", "CON", "SIZ", "DEX", "INT", "POW", "CHA"]

PROFESSIONAL_ROWS = 8

VARIANTS = {
    "standard": {
        "title": "MYTHRAS CHARACTER CREATION SHEET",
        "filename": "Mythras_Character_Sheet_Print.xlsx",
        "steps": ["Culture", "Career"],
        "info": ["Name", "Player", "Culture", "Career", "Age"],
        "d20": False,
        # Fewer columns than d20, so each is wider to use the width of the page.
        "column_width": 7.5,
        "note": "Tick FROM for each step whose skill list offers the skill. "
                "Base: add the two characteristics (×2 doubles one). "
                "Total = base + every point spent on the skill.",
    },
    "d20": {
        "title": "MYTHRAS d20 CLASSIC FANTASY CHARACTER SHEET",
        "filename": "Mythras_d20_Classic_Fantasy_Character_Sheet_Print.xlsx",
        "steps": ["Culture", "Career", "Class"],
        "info": ["Name", "Player", "Culture", "Career", "Class", "Age"],
        "d20": True,
        "column_width": 6.5,
        "note": "Tick FROM for each step whose skill list offers the skill. "
                "Base: add the two characteristics (×2 doubles one). "
                "Total = base + every point spent. "
                "d20 value = total ÷ 5 rounded down; bonus = the remainder.",
    },
}

THIN = Side(style="thin", color="000000")
MEDIUM = Side(style="medium", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNDERLINE = Border(bottom=THIN)

# Light greys only: they survive a black-and-white printer and leave the boxes writable.
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
SECTION_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

TITLE_FONT = Font(bold=True, size=14)
LABEL_FONT = Font(bold=True, size=10)
HEADER_FONT = Font(bold=True, size=8)
BODY_FONT = Font(size=10)
SMALL_FONT = Font(size=9)
NOTE_FONT = Font(italic=True, size=8)

CENTER = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
ON_LINE = Alignment(horizontal="left", vertical="bottom")
RIGHT = Alignment(horizontal="right", vertical="center", indent=1)

SPACER_HEIGHT = 8
ROW_HEIGHT = 18  # tall enough to write in once the page is scaled to fit


def create_print_sheet(variant):
    """Build a blank, print-ready character sheet for one variant."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Character Sheet"
    steps = variant["steps"]

    # Column plan, left to right. Every column after the skill name has the same width, so
    # the characteristic boxes come out equal; the base characteristics are the only field
    # spanning two columns.
    col = 2
    mark_cols = list(range(col, col + len(steps)))
    col += len(steps)
    chars_cols = [col, col + 1]
    base_col = col + 2
    col += 3
    point_cols = list(range(col, col + len(steps) + 1))  # one per step, then free points
    col += len(point_cols)
    total_col = col
    d20_cols = [col + 1, col + 2] if variant["d20"] else []
    last_col = d20_cols[-1] if d20_cols else total_col

    ws.column_dimensions["A"].width = 20
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = variant["column_width"]

    def block(r1, c1, r2, c2, value=None, font=None, alignment=None, border=None, fill=None):
        """Style a rectangle, merging it if it is more than one cell.

        Excel draws a merged range's outline from each member cell's own edges, so borders
        and fills go on every cell, not just the top-left one that holds the value.
        """
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if border:
                    ws.cell(row=r, column=c).border = border
                if fill:
                    ws.cell(row=r, column=c).fill = fill
        top_left = ws.cell(row=r1, column=c1)
        if value is not None:
            top_left.value = value
        if font:
            top_left.font = font
        if alignment:
            top_left.alignment = alignment

    def spacer(row):
        ws.row_dimensions[row].height = SPACER_HEIGHT

    # Title
    row = 1
    block(row, 1, row, last_col, variant["title"], font=TITLE_FONT,
          alignment=Alignment(horizontal="center", vertical="center"),
          border=Border(bottom=MEDIUM))
    ws.row_dimensions[row].height = 26
    spacer(row + 1)

    # Character info: two fields to a row, each a label sitting on a line to write on
    row = 3
    left_last = 1 + (last_col - 2) // 2
    fields = variant["info"]
    for i in range(0, len(fields), 2):
        block(row, 1, row, 1, fields[i], font=LABEL_FONT, alignment=ON_LINE, border=UNDERLINE)
        block(row, 2, row, left_last, border=UNDERLINE)
        if i + 1 < len(fields):
            block(row, left_last + 1, row, left_last + 2, fields[i + 1], font=LABEL_FONT,
                  alignment=Alignment(horizontal="left", vertical="bottom", indent=1),
                  border=UNDERLINE)
            block(row, left_last + 3, row, last_col, border=UNDERLINE)
        ws.row_dimensions[row].height = 22
        row += 1
    spacer(row)

    # Characteristics: a row of labels over a row of boxes
    row += 1
    block(row, 1, row + 1, 1, "CHARACTERISTICS", font=LABEL_FONT,
          alignment=Alignment(horizontal="left", vertical="center"))
    for offset, name in enumerate(CHARACTERISTICS):
        block(row, 2 + offset, row, 2 + offset, name, font=HEADER_FONT, alignment=CENTER,
              border=BOX, fill=HEADER_FILL)
        block(row + 1, 2 + offset, row + 1, 2 + offset, border=BOX)
    ws.row_dimensions[row + 1].height = 26
    row += 2
    spacer(row)

    # Skill table header: groups on the first row, one label per column on the second
    head, sub = row + 1, row + 2

    def header(r1, c1, r2, c2, text):
        block(r1, c1, r2, c2, text, font=HEADER_FONT, alignment=CENTER, border=BOX,
              fill=HEADER_FILL)

    header(head, 1, sub, 1, "SKILL")
    header(head, mark_cols[0], head, mark_cols[-1], "FROM")
    header(head, chars_cols[0], head, base_col, "BASE")
    header(head, point_cols[0], head, point_cols[-1], "POINTS SPENT")
    header(head, total_col, sub, total_col, "TOTAL")
    for c, step in zip(mark_cols, steps):
        header(sub, c, sub, c, step)
    header(sub, chars_cols[0], sub, chars_cols[1], "Chars")
    header(sub, base_col, sub, base_col, "Value")
    for c, step in zip(point_cols, steps + ["Free"]):
        header(sub, c, sub, c, step)
    if d20_cols:
        header(head, d20_cols[0], head, d20_cols[1], "d20")
        header(sub, d20_cols[0], sub, d20_cols[0], "Value")
        header(sub, d20_cols[1], sub, d20_cols[1], "Bonus")
    ws.row_dimensions[head].height = 14
    ws.row_dimensions[sub].height = 14
    row = sub + 1

    def section(row, title):
        block(row, 1, row, last_col, title, font=Font(bold=True, size=9), alignment=LEFT,
              border=BOX, fill=SECTION_FILL)
        ws.row_dimensions[row].height = 15

    def skill_row(row, name=None, chars=None):
        block(row, 1, row, 1, name, font=BODY_FONT, alignment=LEFT, border=BOX)
        block(row, chars_cols[0], row, chars_cols[1], chars, font=SMALL_FONT, alignment=CENTER,
              border=BOX)
        for c in mark_cols + [base_col] + point_cols + [total_col] + d20_cols:
            block(row, c, row, c, border=BOX)
        ws.row_dimensions[row].height = ROW_HEIGHT

    section(row, "STANDARD SKILLS")
    row += 1
    for name, chars in STANDARD_SKILLS:
        skill_row(row, name, chars)
        row += 1

    # Professional skills are chosen per character, so these rows are wholly blank,
    # including the characteristics their base comes from.
    section(row, "PROFESSIONAL SKILLS")
    row += 1
    for _ in range(PROFESSIONAL_ROWS):
        skill_row(row)
        row += 1

    # Pool bookkeeping sits directly under the columns it adds up, so each step's sum is
    # done down its own column.
    for label in ("Points available", "Points spent", "Points left"):
        block(row, 1, row, base_col, label, font=LABEL_FONT, alignment=RIGHT)
        for c in point_cols:
            block(row, c, row, c, border=BOX)
        ws.row_dimensions[row].height = ROW_HEIGHT
        row += 1

    spacer(row)
    row += 1
    block(row, 1, row, last_col, variant["note"], font=NOTE_FONT,
          alignment=Alignment(horizontal="left", vertical="top", wrap_text=True))
    ws.row_dimensions[row].height = 24

    # Page setup: A4 portrait, scaled down to exactly one page, gridlines off so the screen
    # shows what the printer will.
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4, header=0, footer=0)
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{get_column_letter(last_col)}{row}"
    ws.sheet_view.showGridLines = False

    return wb


def main():
    parser = argparse.ArgumentParser(
        description="Write blank Mythras character creation sheets to print and fill in by hand.")
    parser.add_argument("variant", nargs="?", choices=[*VARIANTS, "both"], default="both",
                        help="which sheet to write (default: both)")
    parser.add_argument("-o", "--output-dir", type=Path, default=Path("."),
                        help="directory to write into (default: current directory)")
    args = parser.parse_args()

    names = list(VARIANTS) if args.variant == "both" else [args.variant]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name in names:
        path = args.output_dir / VARIANTS[name]["filename"]
        create_print_sheet(VARIANTS[name]).save(path)
        print(f"Saved {path}")


if __name__ == "__main__":
    main()
